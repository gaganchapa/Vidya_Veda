import streamlit as st
import streamlit.components.v1 as components
from pptx import Presentation
from gtts import gTTS
from langchain.schema import HumanMessage
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.prompts import PromptTemplate
from langchain.chains import LLMChain
import os
import pygame
import time
import json
import pathlib
import hive_final  # Import the PPT generation module

# For Google Drive integration
# import google_upload
import os
import time
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload

# Streamlit page configuration
st.set_page_config(layout="wide", page_title="AI Presentation Generator")

# Initialize the language model
os.environ["GOOGLE_API_KEY"] = "AIzaSyCcf3s3GS7_925D7t2fgODc5WIKOMZSOzc"
llm = ChatGoogleGenerativeAI(
    model="gemini-1.5-pro",
    temperature=0,
    max_tokens=None,
    timeout=None,
    max_retries=2,
)

# Initialize Pygame for audio playback
pygame.mixer.init()

# Initialize session state variables
if 'current_slide' not in st.session_state:
    st.session_state.current_slide = 0
if 'is_playing' not in st.session_state:
    st.session_state.is_playing = False
if 'script_generated' not in st.session_state:
    st.session_state.script_generated = False
if 'audio_files' not in st.session_state:
    st.session_state.audio_files = []
if 'presentation_link' not in st.session_state:
    st.session_state.presentation_link = None
if 'ppt_generated' not in st.session_state:
    st.session_state.ppt_generated = False
if 'slides_text' not in st.session_state:
    st.session_state.slides_text = []
if 'json_data' not in st.session_state:
    st.session_state.json_data = None
if 'output_path' not in st.session_state:
    st.session_state.output_path = None
if 'topic' not in st.session_state:
    st.session_state.topic = ""

# Function to extract text from PPT
def extract_text_from_ppt(ppt_file_path):
    prs = Presentation(ppt_file_path)
    slides_text = []
    for idx, slide in enumerate(prs.slides):
        slide_text = "\n".join([shape.text for shape in slide.shapes if hasattr(shape, "text")])
        slides_text.append(f"Slide {idx + 1}:\n{slide_text}")
    return slides_text

# Function to generate script using LLM
def generate_script(slides_text):
    combined_text = "\n\n---\n\n".join(slides_text)
    prompt = f"""
    You are a teacher explaining a PowerPoint presentation. 
    Generate a detailed, engaging explanation for each slide. 
    Separate each slide's explanation with '---' to indicate pauses.
    
    Slides:
    {combined_text}
    """
    response = llm([HumanMessage(content=prompt)])
    return response.content

# Function to save script as audio
def save_script_as_audio(script):
    slides = script.split("\n\n---\n\n")
    audio_files = []

    for i, slide_text in enumerate(slides):
        tts = gTTS(text=slide_text, lang='en')
        audio_path = f"slide_{i+1}.mp3"
        tts.save(audio_path)
        audio_files.append(audio_path)
    
    return audio_files

# Function to play current slide audio
def play_current_slide():
    if st.session_state.current_slide < len(st.session_state.audio_files):
        pygame.mixer.music.load(st.session_state.audio_files[st.session_state.current_slide])
        pygame.mixer.music.play()
        
        while pygame.mixer.music.get_busy():
            time.sleep(0.1)
            if st.session_state.get("stop_reading"):
                pygame.mixer.music.stop()
                return

# Google Drive upload and publish functions
def authenticate():
    """Authenticate and return a Google Drive service instance."""
    creds = None
    if os.path.exists('token.json'):
        creds = Credentials.from_authorized_user_file('token.json', SCOPES)
    if not creds or not creds.valid:
        flow = InstalledAppFlow.from_client_secrets_file('cred.json', SCOPES)
        creds = flow.run_local_server(port=0)
        with open('token.json', 'w') as token:
            token.write(creds.to_json())

    drive_service = build('drive', 'v3', credentials=creds)
    
    return drive_service

def upload_and_convert_pptx(drive_service, file_path, file_name):
    """Uploads a PowerPoint file to Google Drive and converts it to Google Slides."""
    
    # Upload PPTX to Drive
    file_metadata = {'name': file_name}
    media = MediaFileUpload(file_path, mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    uploaded_file = drive_service.files().create(body=file_metadata, media_body=media, fields='id').execute()
    
    file_id = uploaded_file.get('id')

    # Convert to Google Slides
    converted_file_metadata = {
        'name': file_name,
        'mimeType': 'application/vnd.google-apps.presentation'
    }
    converted_file = drive_service.files().copy(fileId=file_id, body=converted_file_metadata, fields='id').execute()
    converted_file_id = converted_file.get('id')

    print(f"Converted to Google Slides: https://docs.google.com/presentation/d/{converted_file_id}/edit")

    return converted_file_id

def make_public_and_publish(drive_service, file_id):
    """Makes the file public AND publishes it to the web."""
    
    # Step 1: Set file permissions to public
    permission = {'type': 'anyone', 'role': 'reader'}
    drive_service.permissions().create(fileId=file_id, body=permission).execute()
    
    # Step 2: Publish the presentation to the web using Drive API's revisions resource
    try:
        # Get the latest revision
        revisions = drive_service.revisions().list(fileId=file_id).execute()
        revision_id = revisions.get('revisions', [])[-1].get('id')
        
        # Update the revision to make it published
        revision_metadata = {
            'published': True,
            'publishAuto': True,
            'publishedOutsideDomain': True
        }
        
        drive_service.revisions().update(
            fileId=file_id,
            revisionId=revision_id,
            body=revision_metadata
        ).execute()
        
        # Create proper publish URL with parameters
        publish_url = f"https://docs.google.com/presentation/d/{file_id}/pub?start=true&loop=true&delayms=3000"
        embed_url = f"https://docs.google.com/presentation/d/{file_id}/embed?start=true&loop=true&delayms=3000"
        
        return embed_url  # Return the embed URL for iframe
    except Exception as e:
        st.error(f"Error publishing presentation: {e}")
        return None

# Generate JSON content using LLM
def generate_presentation_json(title):
    # Define prompt template
    prompt_template = """
    You are tasked with generating structured JSON content for a PowerPoint presentation on the topic of "{title}". Follow these rules carefully:

    Each slide should contain:
       - "heading": A slide heading that concisely summarizes the main topic of the slide.
       - "bullet_points": 3 to 5 points explaining the content, each point under 50 words, using concise and informative language.
       - "key_message": A one-sentence takeaway that is short and informative, capturing the main idea of the slide.
       - "img_keywords": Maximum of 3 keywords separated by commas, which describe images representing the content. This applies only to slides with regular content (not for Step-by-Step Process or Table slides).
       - For slides comparing features or attributes, include "table_data" as a list of lists in the format [["Feature", "Attribute"], ["Feature A", "Value A1"], ...] to structure comparisons.
       - For slides detailing a process, add a "Step-by-Step Process" under "bullet_points" with points prefixed by '>>'. Each point should be exactly three words, representing each step concisely without including step numbers.

    Include a concluding slide titled "Conclusion" summarizing the main ideas of the presentation.

    Output the content in structured JSON format with these keys for each slide:
      - "heading", "bullet_points", "key_message", "img_keywords"
      - Include "table_data" only when comparing features, and "Step-by-Step Process" only when outlining a process.

    Example JSON format:
    {{
        "title": "{title}",
        "slides": [
            {{
                "heading": "Introduction to {title}",
                "bullet_points": ["...", "..."],
                "key_message": "...",
                "img_keywords": "keyword1,keyword2,keyword3"
            }},
            {{
                "heading": "{title} in Everyday Life",
                "bullet_points": ["...", "..."],
                "key_message": "...",
                "img_keywords": "keyword1,keyword2,keyword3"
            }},
            {{
                "heading": "Comparison of {title} Features",
                "bullet_points": [],
                "table_data": [["Feature", "Attribute"], ["Feature A", "Value A1"]],
                "key_message": "...",
                "img_keywords": null
            }},
            {{
                "heading": "Step-by-Step Process for {title}",
                "bullet_points": [">> Define Problem", ">> Gather Data", ">> Model Training"],
                "key_message": "...",
                "img_keywords": null
            }},
            {{
                "heading": "Conclusion: Embracing {title}'s Potential",
                "bullet_points": ["...", "..."],
                "key_message": "...",
                "img_keywords": "keyword1,keyword2,keyword3"
            }}
        ]
    }}
    """

    # Create the chain with the prompt and LLM
    prompt = PromptTemplate(
        input_variables=["title"],
        template=prompt_template
    )
    chain = LLMChain(llm=llm, prompt=prompt)
    
    # Run the chain to generate content
    output = chain.run(title=title)
    
    # Clean up the output and convert to JSON
    out = output.replace("```json", "").replace("```", "")
    
    try:
        json_output = json.loads(out)
        with open('output.json', 'w') as json_file:
            json.dump(json_output, json_file, indent=4)
        return json_output
    except json.JSONDecodeError as e:
        st.error(f"Error parsing JSON: {e}")
        return None

# Function to add a new slide to existing presentation JSON
def generate_new_slide(json_data, slide_position, slide_topic):
    # Define the prompt for generating a new slide
    prompt_template = """
    You are tasked with generating a single slide for a PowerPoint presentation to be added at position {position} on the topic of "{topic}". 
    The slide should integrate well with the existing presentation titled "{title}".

    The slide should contain:
       - "heading": A slide heading that concisely summarizes the main topic of the slide.
       - "bullet_points": 3 to 5 points explaining the content, each point under 30 words, using concise and informative language.
       - "key_message": A one-sentence takeaway that is short and informative, capturing the main idea of the slide.
       - "img_keywords": Maximum of 3 keywords separated by commas, which describe images representing the content.

    Output only the JSON object for this single slide:
    {{
        "heading": "...",
        "bullet_points": ["...", "...", "..."],
        "key_message": "...",
        "img_keywords": "keyword1,keyword2,keyword3"
    }}
    """

    # Create the chain with the prompt and LLM
    prompt = PromptTemplate(
        input_variables=["position", "topic", "title"],
        template=prompt_template
    )
    chain = LLMChain(llm=llm, prompt=prompt)
    
    # Run the chain to generate content
    output = chain.run(position=slide_position, topic=slide_topic, title=json_data["title"])
    
    # Clean up the output and convert to JSON
    out = output.replace("```json", "").replace("```", "")
    
    try:
        slide_json = json.loads(out)
        return slide_json
    except json.JSONDecodeError as e:
        st.error(f"Error parsing JSON for new slide: {e}")
        return None

# Function to edit an existing slide
def edit_slide(json_data, slide_index, updated_heading, updated_bullet_points, updated_key_message):
    try:
        # Get the existing slide
        slide = json_data["slides"][slide_index]
        
        # Update slide values if provided
        if updated_heading:
            slide["heading"] = updated_heading
        
        if updated_bullet_points:
            # Parse bullet points from text area (assuming one bullet point per line)
            bullet_points_list = [bp.strip() for bp in updated_bullet_points.split('\n') if bp.strip()]
            slide["bullet_points"] = bullet_points_list
        
        if updated_key_message:
            slide["key_message"] = updated_key_message
        
        # Update the slide in the JSON data
        json_data["slides"][slide_index] = slide
        
        return json_data
    except Exception as e:
        st.error(f"Error editing slide: {e}")
        return None

# Function to regenerate PPT and update Google Drive link
def regenerate_and_publish_ppt(json_data, topic):
    # Save updated JSON
    with open('output.json', 'w') as json_file:
        json.dump(json_data, json_file, indent=4)
    
    # Generate PowerPoint presentation
    output_path = pathlib.Path('output_presentation_2.pptx')
    
    try:
        hive_final.generate_powerpoint_presentation(json_data, output_path)
        st.session_state.ppt_generated = True
        st.session_state.output_path = output_path
        
        # Extract text from the presentation for script generation
        st.session_state.slides_text = extract_text_from_ppt(output_path)
        
        # Upload to Google Drive and publish
        drive_service = authenticate()
        file_id = upload_and_convert_pptx(drive_service, str(output_path), f"{topic} Presentation")
        
        # Wait for Google to process
        time.sleep(10)
        
        # Publish the presentation
        embed_url = make_public_and_publish(drive_service, file_id)
        
        if embed_url:
            st.session_state.presentation_link = embed_url
            return True
        else:
            st.error("Failed to publish presentation.")
            return False
    except Exception as e:
        st.error(f"Error regenerating PowerPoint: {e}")
        return False

# Define Google API scope
SCOPES = ['https://www.googleapis.com/auth/drive']

# -------------------- Streamlit UI --------------------
st.title("AI Presentation Generator")

# Step 1: Enter presentation topic
topic = st.text_input("Enter the topic for your presentation:", 
                     placeholder="e.g., AI Applications in Healthcare",
                     value=st.session_state.topic)

if topic != st.session_state.topic:
    st.session_state.topic = topic

# Generate presentation button
if st.button("Generate Presentation") and topic:
    with st.spinner("Generating presentation content..."):
        # Generate JSON content
        json_data = generate_presentation_json(topic)
        
        if json_data:
            # Save JSON data to session state
            st.session_state.json_data = json_data
            
            # Save to output.json
            with open('output.json', 'w') as json_file:
                json.dump(json_data, json_file, indent=4)
            
            # Generate PowerPoint presentation
            output_path = pathlib.Path('output_presentation_2.pptx')
            st.session_state.output_path = output_path
            
            try:
                hive_final.generate_powerpoint_presentation(json_data, output_path)
                st.session_state.ppt_generated = True
                st.success("PowerPoint presentation generated successfully!")
                
                # Extract text from the presentation for script generation
                st.session_state.slides_text = extract_text_from_ppt(output_path)
                
                # Upload to Google Drive and publish
                with st.spinner("Uploading presentation to Google Drive..."):
                    try:
                        drive_service = authenticate()
                        file_id = upload_and_convert_pptx(drive_service, str(output_path), f"{topic} Presentation")
                        
                        # Wait for Google to process
                        time.sleep(10)
                        
                        # Publish the presentation
                        embed_url = make_public_and_publish(drive_service, file_id)
                        
                        if embed_url:
                            st.session_state.presentation_link = embed_url
                            st.success("Presentation uploaded and published successfully!")
                        else:
                            st.error("Failed to publish presentation.")
                    except Exception as e:
                        st.error(f"Error in Google Drive integration: {e}")
            except Exception as e:
                st.error(f"Error generating PowerPoint: {e}")
        else:
            st.error("Failed to generate presentation content.")

# Display the presentation if available
if st.session_state.presentation_link:
    st.write("### Your Presentation")
    components.iframe(st.session_state.presentation_link, height=600, width=960)
    
    # Slide modifications section
    if st.session_state.json_data:
        # Create tabs for different modification options
        tab1, tab2 = st.tabs(["Add New Slide", "Edit Existing Slide"])
        
        # Tab 1: Add New Slide
        with tab1:
            st.write("### Add a New Slide")
            
            # Calculate number of slides
            num_slides = len(st.session_state.json_data["slides"])
            
            col1, col2 = st.columns(2)
            
            with col1:
                slide_position = st.number_input("Insert slide at position", 
                                               min_value=1, 
                                               max_value=num_slides+1, 
                                               value=num_slides+1,
                                               help="Position 1 is the first slide. Choose a position up to one more than the current number of slides.")
            
            with col2:
                slide_topic = st.text_input("New slide topic", 
                                          placeholder="e.g., Benefits of AI in Healthcare",
                                          help="Enter a specific topic for this slide that relates to the main presentation topic.")
            
            slide_content = st.text_area("Custom notes for slide content (optional)", 
                                       placeholder="Add any specific points or information you want to include in this slide.",
                                       help="This will help guide the AI in generating more relevant slide content.")
            
            if st.button("Add Slide") and slide_topic:
                with st.spinner("Generating new slide and updating presentation..."):
                    # Generate new slide content
                    new_slide = generate_new_slide(st.session_state.json_data, slide_position, slide_topic)
                    
                    if new_slide:
                        # Add additional context if provided
                        if slide_content:
                            # Modify the prompt for the new slide to include custom content
                            prompt = f"""
                            Refine the bullet points for this slide on "{slide_topic}" to incorporate these specific notes:
                            {slide_content}
                            
                            Return only the updated bullet points as a JSON list:
                            """
                            response = llm([HumanMessage(content=prompt)])
                            try:
                                updated_points = json.loads(response.content.replace("```json", "").replace("```", ""))
                                new_slide["bullet_points"] = updated_points
                            except:
                                st.warning("Could not parse custom content. Using AI-generated content instead.")
                        
                        # Insert the new slide at the specified position
                        position_idx = int(slide_position) - 1
                        st.session_state.json_data["slides"].insert(position_idx, new_slide)
                        
                        # Update the presentation
                        success = regenerate_and_publish_ppt(st.session_state.json_data, topic)
                        
                        if success:
                            st.success(f"New slide on '{slide_topic}' added successfully and presentation updated!")
                            # Display the updated iframe directly
                            st.write("### Updated Presentation")
                            components.iframe(st.session_state.presentation_link, height=600, width=960)
                        else:
                            st.error("Failed to update presentation with new slide.")
                    else:
                        st.error("Failed to generate new slide content.")
        
        # Tab 2: Edit Existing Slide
        with tab2:
            st.write("### Edit an Existing Slide")
            
            # Calculate number of slides
            num_slides = len(st.session_state.json_data["slides"])
            
            # Select slide to edit
            slide_to_edit = st.number_input("Select slide to edit", 
                                           min_value=1, 
                                           max_value=num_slides, 
                                           value=1,
                                           help="Choose the slide number you want to edit.")
            
            # Get the selected slide
            slide_index = int(slide_to_edit) - 1
            current_slide = st.session_state.json_data["slides"][slide_index]
            
            # Show current slide content
            st.write("#### Current Slide Content")
            st.write(f"**Heading:** {current_slide['heading']}")
            st.write("**Bullet Points:**")
            for point in current_slide.get("bullet_points", []):
                st.write(f"- {point}")
            st.write(f"**Key Message:** {current_slide.get('key_message', '')}")
            
            # Edit fields
            st.write("#### Edit Slide Content")
            new_heading = st.text_input("Update heading", 
                                      value=current_slide['heading'],
                                      help="Edit the slide heading")
            
            current_bullet_points = "\n".join(current_slide.get("bullet_points", []))
            new_bullet_points = st.text_area("Update bullet points (one per line)", 
                                           value=current_bullet_points,
                                           help="Edit the bullet points, one per line")
            
            new_key_message = st.text_input("Update key message", 
                                          value=current_slide.get('key_message', ''),
                                          help="Edit the key message of the slide")
            
            if st.button("Update Slide"):
                with st.spinner("Updating slide and regenerating presentation..."):
                    # Update the slide in the JSON data
                    updated_json_data = edit_slide(
                        st.session_state.json_data,
                        slide_index,
                        new_heading,
                        new_bullet_points,
                        new_key_message
                    )
                    
                    if updated_json_data:
                        st.session_state.json_data = updated_json_data
                        
                        # Regenerate and publish the updated presentation
                        success = regenerate_and_publish_ppt(updated_json_data, topic)
                        
                        if success:
                            st.success(f"Slide {slide_to_edit} updated successfully!")
                            # Display the updated iframe directly
                            st.write("### Updated Presentation")
                            components.iframe(st.session_state.presentation_link, height=600, width=960)
                        else:
                            st.error("Failed to update the presentation.")
                    else:
                        st.error("Failed to update slide content.")
    
    # Generate script and audio section
    if st.button("Generate Explanatory Script & Audio"):
        with st.spinner("Generating script and converting to audio..."):
            script = generate_script(st.session_state.slides_text)
            audio_files = save_script_as_audio(script)
            st.session_state.audio_files = audio_files
            st.session_state.script_generated = True
            st.success("Script and audio generated successfully!")
    
    # Audio player controls
    if st.session_state.script_generated:
        st.write(f"### Current Slide: {st.session_state.current_slide + 1}")
        col1, col2, col3 = st.columns(3)
        
        with col1:
            if st.button("⬅ Previous Slide") and st.session_state.current_slide > 0:
                st.session_state.current_slide -= 1
                play_current_slide()
        
        with col2:
            if st.button("Play/Resume Explanation"):
                st.session_state.is_playing = True
                st.session_state.stop_reading = False
                play_current_slide()
        
        with col3:
            if st.button("Next Slide ➡") and st.session_state.current_slide < len(st.session_state.audio_files) - 1:
                st.session_state.current_slide += 1
                play_current_slide()
        
        if st.button("Stop Explanation"):
            st.session_state.stop_reading = True
            pygame.mixer.music.stop()
            st.session_state.is_playing = False
else:
    if not topic:
        st.info("Enter a presentation topic above and click 'Generate Presentation' to start.")
    elif not st.session_state.ppt_generated:
        st.info("Once you generate a presentation, it will appear here.")